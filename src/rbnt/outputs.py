"""The non-web editions: PDF, chart PNGs, Word file, workbook and social cards.

The web page, the PDF and the Word file are all made from the same content
blocks, so the three editions cannot drift apart.
"""

import csv
import os
import re
from html import unescape
from pathlib import Path

from . import edition
from .drawing import STYLE

HOME = edition.HOME
TAGS = re.compile(r"<[^>]+>")


def text_of(html):
    return unescape(TAGS.sub("", html))


def browser(pw):
    """Chromium: CHROMIUM_PATH if set, else Playwright's own, else the system one under /opt/pw-browsers."""
    if os.environ.get("CHROMIUM_PATH"):
        return pw.chromium.launch(executable_path=os.environ["CHROMIUM_PATH"])
    try:
        return pw.chromium.launch()
    except Exception:
        return pw.chromium.launch(executable_path="/opt/pw-browsers/chromium")


def font_faces(prefix="BVP", mono="JBM"):
    folder = (HOME / "assets" / "fonts").as_uri()
    faces = "".join(f'@font-face{{font-family:"{prefix}";font-weight:{w};src:url("{folder}/be-vietnam-pro-{w}.woff2")}}'
                    for w in (400, 500, 600, 700))
    return faces + f'@font-face{{font-family:"{mono}";font-weight:400 600;src:url("{folder}/jetbrains-mono.woff2")}}'


def print_pdf(page, target, footer):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        b = browser(pw)
        tab = b.new_page()
        tab.goto(Path(page).as_uri(), wait_until="networkidle")
        tab.evaluate("document.documentElement.dataset.mode = 'light'")
        tab.emulate_media(media="print")
        tab.wait_for_timeout(300)
        tab.pdf(path=str(target), format="A4", print_background=True, display_header_footer=True,
                header_template="<span></span>",
                footer_template=('<div style="width:100%;padding:0 13mm;display:flex;justify-content:space-between;'
                                 f'font:7px monospace;color:#596975"><span>{footer}</span>'
                                 '<span><span class="pageNumber"></span> / <span class="totalPages"></span></span></div>'),
                margin={"top": "14mm", "bottom": "16mm", "left": "13mm", "right": "13mm"})
        b.close()


def svg_to_png(names, folder, scale=2):
    from playwright.sync_api import sync_playwright
    folder.mkdir(parents=True, exist_ok=True)
    faces = font_faces("Be Vietnam Pro", "JetBrains Mono")
    with sync_playwright() as pw:
        b = browser(pw)
        tab = b.new_page(device_scale_factor=scale)
        for name in names:
            svg = (HOME / "figures" / f"{name}.svg").read_text()
            tab.set_content(f"<html><head><style>{faces}body{{margin:0;background:#fff}}{STYLE}</style></head>"
                            f"<body><div id=frame style='display:inline-block;padding:8px'>{svg}</div></body></html>")
            tab.wait_for_timeout(120)
            tab.locator("#frame").screenshot(path=str(folder / f"{name}.png"))
        b.close()


def word(blocks, target, png_folder, top, meta):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    ink, ink2, faint, red = (RGBColor(0x0F, 0x1A, 0x22), RGBColor(0x39, 0x48, 0x53), RGBColor(0x59, 0x69, 0x75),
                             RGBColor(0xB3, 0x27, 0x2A))
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(1.8)
    base = doc.styles["Normal"]
    base.font.name = "Be Vietnam Pro"
    base.element.rPr.rFonts.set(qn("w:eastAsia"), "Be Vietnam Pro")
    base.font.size = Pt(10.5)
    base.font.color.rgb = ink2
    base.paragraph_format.line_spacing = 1.35
    base.paragraph_format.space_after = Pt(6)
    for style, size in (("Heading 1", 22), ("Heading 2", 16), ("Heading 3", 12.5)):
        h = doc.styles[style]
        h.font.name, h.font.size, h.font.bold, h.font.color.rgb = "Be Vietnam Pro", Pt(size), True, ink

    def fill(cell, colour):
        props = cell._tc.get_or_add_tcPr()
        shade = OxmlElement("w:shd")
        shade.set(qn("w:val"), "clear")
        shade.set(qn("w:fill"), colour)
        props.append(shade)

    def runs(par, html, size=None, colour=None):
        for piece in re.split(r"(<b>.*?</b>|<strong>.*?</strong>)", html):
            if piece:
                r = par.add_run(text_of(piece))
                r.bold = piece.startswith(("<b>", "<strong>"))
                if r.bold:
                    r.font.color.rgb = ink
                if size:
                    r.font.size = Pt(size)
                if colour and not r.bold:
                    r.font.color.rgb = colour

    def small_caps(text, colour=faint):
        r = doc.add_paragraph().add_run(text.upper())
        r.bold, r.font.size, r.font.color.rgb = True, Pt(8), colour

    def grid(head, rows, mono=None):
        mono = set(range(1, len(head))) if mono is None else mono
        t = doc.add_table(rows=1, cols=len(head))
        t.style, t.alignment = "Table Grid", WD_TABLE_ALIGNMENT.CENTER
        for k, h in enumerate(head):
            c = t.rows[0].cells[k]
            c.text = ""
            r = c.paragraphs[0].add_run(text_of(str(h)))
            r.bold, r.font.size, r.font.color.rgb = True, Pt(7.5), faint
            fill(c, "E6EBEF")
        for row in rows:
            cells = t.add_row().cells
            for k, v in enumerate(row):
                cells[k].text = ""
                p = cells[k].paragraphs[0]
                r = p.add_run(text_of(str(v).replace("<br>", " ")))
                r.font.size = Pt(8)
                if k in mono:
                    r.font.name = "JetBrains Mono"
                    p.alignment = 2
        doc.add_paragraph()

    def boxed(paragraph_writer):
        t = doc.add_table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        fill(cell, "F4F6F8")
        cell.text = ""
        paragraph_writer(cell)
        doc.add_paragraph()

    logo = HOME / "brand" / "logos" / "dao-logo-on-light.png"
    if logo.exists():
        doc.add_picture(str(logo), height=Cm(2))
    small_caps(top["overline"], red)
    doc.add_heading(top["title"], 1)
    runs(doc.add_paragraph(), top["sub"])
    small_caps("The quarter in numbers")
    grid(["", *top["board_head"], "Change"], top["board"])
    for b in blocks:
        kind = b[0]
        if kind == "chapter":
            doc.add_page_break()
            small_caps(f"Chapter {b[1]}", red)
            doc.add_heading(b[2], 2)
        elif kind == "lead":
            runs(doc.add_paragraph(), b[1], 12, ink)
        elif kind == "text":
            runs(doc.add_paragraph(), b[1])
        elif kind == "h3":
            doc.add_heading(b[1], 3)
        elif kind == "bullets":
            for item in b[1]:
                runs(doc.add_paragraph(style="List Bullet"), item)
        elif kind == "findings":
            for value, title, desc in b[1]:
                p = doc.add_paragraph()
                lead = p.add_run(f"{value}  {title}. ")
                lead.bold, lead.font.color.rgb = True, ink
                runs(p, desc)
        elif kind == "rules":
            for title, desc in b[1]:
                p = doc.add_paragraph(style="List Bullet")
                runs(p, f"<b>{title}.</b> {desc}")
        elif kind in ("terms", "aside"):
            def write(cell, b=b):
                label = "Three terms used throughout" if b[0] == "terms" else b[1]
                r = cell.paragraphs[0].add_run(label.upper())
                r.bold, r.font.size, r.font.color.rgb = True, Pt(8), faint
                if b[0] == "terms":
                    for t_, d in b[1]:
                        runs(cell.add_paragraph(), f"<b>{t_}.</b> {d}")
                else:
                    runs(cell.add_paragraph(), b[2])
            boxed(write)
        elif kind == "chart":
            _, number, title, note, name, provenance, _csv, _grid = b
            small_caps(f"Chart {number}")
            r = doc.add_paragraph().add_run(title)
            r.bold, r.font.color.rgb = True, ink
            runs(doc.add_paragraph(), note, 9)
            if (png_folder / f"{name}.png").exists():
                doc.add_picture(str(png_folder / f"{name}.png"), width=Cm(16.5))
            runs(doc.add_paragraph(), f"Read from {provenance}", 7.5, faint)
        elif kind == "grid":
            grid(*b[1:])
        elif kind == "raw" and len(b) > 2:
            grid(*b[2][1:])
        elif kind == "tryit":
            runs(doc.add_paragraph(), b[1])
            r = doc.add_paragraph().add_run(b[2])
            r.font.name, r.font.size = "JetBrains Mono", Pt(7.5)
        elif kind == "files":
            for title, kind_, href in b[1]:
                runs(doc.add_paragraph(style="List Bullet"), f"<b>{title}</b>: {href}")
        elif kind == "glossary":
            for term, desc in b[1]:
                runs(doc.add_paragraph(), f"<b>{term}.</b> {desc}")
    doc.core_properties.title = meta["title"]
    doc.core_properties.author = "Redbelly DAO"
    doc.save(str(target))


def workbook(target, tables, about):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    front = wb.active
    front.title = "Read me"
    for row in about:
        front.append(row)
    front["A1"].font = Font(bold=True, size=14)
    front.column_dimensions["A"].width, front.column_dimensions["B"].width = 26, 110
    for row in front.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    shade = PatternFill("solid", fgColor="E6EBEF")
    for name in tables:
        sheet = wb.create_sheet(name.removesuffix(".csv")[:31])
        with open(edition.out_dir() / name) as fh:
            for row in csv.reader(fh):
                sheet.append([int(v) if re.fullmatch(r"-?\d+", v) else float(v) if re.fullmatch(r"-?\d+\.\d+", v) else v
                              for v in row])
        for cell in sheet[1]:
            cell.font, cell.fill = Font(bold=True), shade
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(48, max(12, max(len(str(c.value or "")) for c in col) + 2))
            for cell in col[1:]:
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00" if abs(cell.value) >= 1000 else "0.0000"
                elif isinstance(cell.value, int):
                    cell.number_format = "#,##0"
        sheet.freeze_panes = "B2"
    wb.properties.creator = "Redbelly DAO"
    wb.properties.title = about[0][0]
    wb.save(str(target))


def cards(cfg, both, c1, block, read_on):
    """community/card-x.png (1600x900) and community/card-discord.png (1200x630)."""
    from playwright.sync_api import sync_playwright
    logo = (HOME / "brand" / "logos" / "dao-logo-on-dark.png").as_uri()
    sign = "+" if both["change_pct"] > 0 else ""
    facts = [(f"{both['rbnt_end'] / 1e6:,.1f}M", "RBNT staked in Reddex pools", f"{sign}{both['change_pct']:.1f}% on 30 June"),
             (f"{c1['wallets']:,}", "wallets holding 1 RBNT or more", "on 30 September"),
             (f"{c1['top10']:.1f}%", "of wallet RBNT in the top 10", "on 30 September")]
    for w, h, name, k in ((1600, 900, "card-x", 1.0), (1200, 630, "card-discord", 0.75)):
        rows = "".join(f'<li><b>{v}</b><span>{t}</span><i>{d}</i></li>' for v, t, d in facts)
        page = f"""<html><head><style>{font_faces()}
        *{{box-sizing:border-box;margin:0}}body{{width:{w}px;height:{h}px;background:#0c1317;color:#f1f5f8;font-family:BVP;
        display:grid;grid-template-columns:1.1fr 1fr;gap:{56 * k}px;padding:{76 * k}px;overflow:hidden}}
        .l{{display:flex;flex-direction:column}} img{{height:{64 * k}px;width:auto;align-self:flex-start}}
        .o{{font-family:JBM;font-size:{18 * k}px;color:#ffb3ae;margin-top:{42 * k}px}}
        h1{{font-size:{70 * k}px;line-height:1.04;letter-spacing:-.02em;margin-top:{14 * k}px}}
        .f{{margin-top:auto;font-family:JBM;font-size:{15 * k}px;color:#8fa1ad;line-height:1.6}}
        ul{{list-style:none;padding:0;display:flex;flex-direction:column;justify-content:center;gap:{18 * k}px}}
        li{{background:#141f26;border:1px solid #263641;border-left:{6 * k}px solid #EF5350;border-radius:8px;padding:{22 * k}px {26 * k}px}}
        li b{{display:block;font-size:{56 * k}px;letter-spacing:-.02em}} li span{{display:block;font-size:{20 * k}px;color:#c2ced6}}
        li i{{display:block;font-style:normal;font-family:JBM;font-size:{14 * k}px;color:#8fa1ad;margin-top:{6 * k}px}}
        </style></head><body><div class="l"><img src="{logo}"><p class="o">On-chain holder report · {cfg['edition']}</p>
        <h1>What RBNT holders did, July to September</h1>
        <p class="f">Redbelly mainnet, block {block:,}. Read {read_on}.<br>Groups only. No forecasts, no advice.</p></div>
        <ul>{rows}</ul></body></html>"""
        with sync_playwright() as pw:
            b = browser(pw)
            tab = b.new_page(viewport={"width": w, "height": h})
            tab.set_content(page)
            tab.wait_for_timeout(300)
            tab.screenshot(path=str(HOME / "community" / f"{name}.png"))
            b.close()
